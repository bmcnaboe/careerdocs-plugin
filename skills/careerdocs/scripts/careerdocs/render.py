"""Render a content plan into the DOCX template, and record what was produced.

``render`` fills the template (docxtpl) from the content plan and writes the document under
the application's ``outputs/`` (a baseline under ``baselines/<positioning>/``) as
``<Name>-<Kind>.docx`` — ``outputs.file_name`` in the config, with ``{name}``, ``{kind}``,
``{org}``, and ``{slug}`` placeholders — beside an output-record skeleton carrying the union
of the plan's source ids and every check marked pending. The newest render always carries
the plain name: a previous render of that name is rotated to a ``_bak1`` suffix (``_bak1``
to ``_bak2``, and so on) together with its PDF, record, and layout renders, so nothing is
overwritten or lost. ``--pdf`` converts via LibreOffice (`soffice`) when it is on PATH, and
records the conversion as skipped with a reason when it is not. Every link the profile
holds (contact and project links, patent URLs, the contact email) appears in the text as
its bare display form and is made a real hyperlink after rendering, so the DOCX and the
PDF are clickable without any template placeholder.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
from collections import defaultdict
from pathlib import Path

from . import __version__
from . import config as config_module
from . import plan as plan_module
from . import schema, util
from .errors import CareerDocsError
from .providers import load_provider

_PDF_CHECKS = ("pagination", "layout")
_ALL_CHECKS = ("factual", "links_dates", "extraction", "pagination", "layout")


_KIND_LABELS = {"resume": "Resume", "cover_letter": "Cover-Letter"}
_SIDE_SUFFIXES = (".docx", ".pdf", ".record.json")
_BAK_RE = re.compile(r"^(?P<stem>.+)_bak(?P<n>\d+)$")


def slugify(text: str) -> str:
    """``Jordan Rivera`` → ``Jordan-Rivera``: words joined by hyphens, punctuation dropped."""
    return "-".join(re.findall(r"[^\W_]+", text))


def document_stem(plan: dict, template: dict, kind: str, pattern: str, *, slug: str = "", org: str = "") -> str:
    """The file name (without suffix) from ``outputs.file_name``: ``{name}`` is the contact
    unit's first line, ``{kind}`` Resume or Cover-Letter, ``{org}`` the brief's organization,
    ``{slug}`` the application slug (``baseline-<positioning>`` for a baseline)."""
    contact = ""
    for section in template["sections"]:
        if section.get("placeholder") == "contact" or "contact" in section.get("entity_types", []):
            units = [u for u in plan["units"] if u["section_id"] == section["id"]]
            contact = units[0]["text"] if units else ""
    values = {"name": slugify(contact.partition("\n")[0]), "kind": _KIND_LABELS.get(kind, kind),
              "org": slugify(org), "slug": slugify(slug)}
    try:
        stem = pattern.format(**values)
    except (KeyError, IndexError, ValueError) as exc:
        raise CareerDocsError(f"outputs.file_name {pattern!r} is not a valid pattern: {exc}", code="USAGE") from exc
    stem = re.sub(r"-{2,}", "-", stem).strip("-")
    return stem or values["kind"]


def _move_render(directory: Path, old_stem: str, new_stem: str) -> None:
    for suffix in _SIDE_SUFFIXES:
        source = directory / f"{old_stem}{suffix}"
        if source.exists():
            source.rename(directory / f"{new_stem}{suffix}")
    layout = directory / "layout"
    if layout.is_dir():
        page = re.compile(rf"^{re.escape(old_stem)}(-p\d+\.png)$")
        for png in sorted(layout.iterdir()):
            match = page.match(png.name)
            if match:
                png.rename(layout / f"{new_stem}{match.group(1)}")
    record_path = directory / f"{new_stem}.record.json"
    if record_path.is_file():
        record = json.loads(record_path.read_text(encoding="utf-8"))
        record["document"] = str(directory / f"{new_stem}.docx")
        if record.get("pdf"):
            record["pdf"] = str(directory / f"{new_stem}.pdf")
        record_path.write_text(json.dumps(record, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def rotate_previous(directory: Path, stem: str) -> None:
    """Free ``stem`` for a new render by shifting every earlier render up one backup slot:
    the plain name becomes ``_bak1``, ``_bak1`` becomes ``_bak2``, and so on. The document,
    its PDF, its record (paths rewritten), and its layout renders move together."""
    if not any((directory / f"{stem}{suffix}").exists() for suffix in _SIDE_SUFFIXES):
        return
    slots = {0}
    for path in directory.iterdir():
        match = _BAK_RE.match(path.name.split(".")[0])
        if match and match.group("stem") == stem:
            slots.add(int(match.group("n")))
    for n in sorted(slots, reverse=True):
        _move_render(directory, stem if n == 0 else f"{stem}_bak{n}", f"{stem}_bak{n + 1}")


def link_map(profile: dict) -> dict[str, str]:
    """Display form → target for every link the visible profile holds: entity ``links[]``
    and ``url`` fields as bare domains, and the contact email as a ``mailto:`` link."""
    links: dict[str, str] = {}
    for entity in profile["entities"]:
        if not plan_module.is_visible(entity):
            continue
        for url in [link.get("url") for link in entity.get("links") or []] + [entity.get("url")]:
            if url:
                links[plan_module.display_link(url)] = url
        if entity["type"] == "contact" and entity.get("email"):
            links[entity["email"]] = f"mailto:{entity['email']}"
    return links


def linkify(docx_path: Path, links: dict[str, str]) -> int:
    """Wrap each occurrence of a link's display text in a hyperlink to its target, keeping
    the run's own formatting so the text reads as before. Returns the number of links."""
    if not links:
        return 0
    import copy

    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.run import Run

    document = Document(str(docx_path))
    pattern = re.compile("|".join(re.escape(d) for d in sorted(links, key=len, reverse=True)))
    made = 0
    for paragraph in document.paragraphs:
        for run in list(paragraph.runs):
            text = run.text
            if not pattern.search(text):
                continue
            pieces: list[tuple[str, str | None]] = []
            position = 0
            for match in pattern.finditer(text):
                if match.start() > position:
                    pieces.append((text[position:match.start()], None))
                pieces.append((match.group(0), links[match.group(0)]))
                position = match.end()
            if position < len(text):
                pieces.append((text[position:], None))
            anchor = run._r
            for piece, target in pieces:
                new_run = copy.deepcopy(anchor)
                Run(new_run, paragraph).text = piece
                element = new_run
                if target:
                    hyperlink = OxmlElement("w:hyperlink")
                    hyperlink.set(qn("r:id"), paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                    hyperlink.append(new_run)
                    element = hyperlink
                    made += 1
                anchor.addprevious(element)
            anchor.getparent().remove(anchor)
    if made:
        document.save(str(docx_path))
    return made


def _unit_context(unit: dict) -> dict:
    # The first line splits at a tab into a head and a tail (an experience unit's role and
    # dates); any further lines are the note (a role's summary), so a template can style the
    # three parts — a bold role, a right-aligned date, a plain descriptor line — separately.
    first, _, note = unit["text"].partition("\n")
    head, _, tail = first.partition("\t")
    return {"text": unit["text"], "kind": unit["kind"], "head": head, "tail": tail, "note": note}


def build_context(plan: dict, template: dict) -> dict:
    by_section: dict[str, list] = defaultdict(list)
    for unit in plan["units"]:
        by_section[unit["section_id"]].append(unit)

    contact_text = ""
    sections = []
    for section in template["sections"]:
        units = by_section.get(section["id"], [])
        if section.get("placeholder") == "contact" or "contact" in section.get("entity_types", []):
            contact_text = units[0]["text"] if units else ""
            continue
        sections.append({
            "id": section["id"],
            "title": section.get("title", ""),
            "units": [_unit_context(u) for u in units],
        })
    # The contact unit is two lines: the name, then the details line.
    contact_name, _, contact_details = contact_text.partition("\n")
    return {
        "contact": contact_text,
        "contact_name": contact_name,
        "contact_details": contact_details,
        "positioning": plan["positioning"],
        "sections": sections,
    }


def render_document(plan: dict, template: dict, template_docx: Path, out_path: Path) -> None:
    from docxtpl import DocxTemplate

    doc = DocxTemplate(str(template_docx))
    doc.render(build_context(plan, template))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def convert_to_pdf(docx_path: Path) -> Path | None:
    if shutil.which("soffice") is None:
        return None
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        check=True, capture_output=True, text=True,
    )
    pdf_path = docx_path.with_suffix(".pdf")
    return pdf_path if pdf_path.exists() else None


def build_record(plan: dict, *, document: Path, kind: str, positioning: str,
                 plan_path: Path | None, brief_path: Path | None, map_path: Path | None,
                 pdf_path: Path | None, pdf_available: bool) -> dict:
    source_ids = sorted({sid for unit in plan["units"] for sid in unit["source_ids"]})
    checks = {name: {"status": "pending", "details": ""} for name in _ALL_CHECKS}
    if not pdf_available:
        for name in _PDF_CHECKS:
            checks[name] = {"status": "skipped", "details": "no PDF (soffice not available)"}
    record = {
        "document": str(document),
        "kind": kind,
        "generated_at": util.now(),
        "plugin_version": __version__,
        "schema_version": schema.profile_schema_version(),
        "role_brief": str(brief_path) if brief_path else None,
        "map": str(map_path) if map_path else None,
        "content_plan": str(plan_path) if plan_path else None,
        "template": plan["template"],
        "voice": plan.get("voice"),
        "positioning": positioning,
        "source_ids": source_ids,
        "checks": checks,
        "stale": False,
        "stale_reason": None,
    }
    if pdf_path:
        record["pdf"] = str(pdf_path)
    return record


# --- CLI ---


def register(subparsers, common: argparse.ArgumentParser) -> None:
    parser = subparsers.add_parser("render", parents=[common], help="render a content plan into the template")
    parser.add_argument("--kind", choices=["resume", "cover_letter"], default="resume")
    parser.add_argument("--role-slug")
    parser.add_argument("--positioning", choices=["executive", "builder"], help="baseline positioning")
    parser.add_argument("--baseline", action="store_true", help="render a role-less baseline")
    parser.add_argument("--pdf", action="store_true", help="also produce a PDF (needs soffice)")
    parser.set_defaults(func=cmd_render)


def cmd_render(args) -> int:
    cfg = config_module.resolve_config(args.workspace)
    if args.baseline:
        positioning = args.positioning or cfg["workflow"]["positioning_default"]
        app_dir = Path(args.workspace) / cfg["outputs"]["baselines_dir"] / positioning
        brief_path = map_path = None
    else:
        slug = plan_module._resolve_slug(args, cfg)
        app_dir = Path(args.workspace) / cfg["outputs"]["applications_dir"] / slug
        brief_path = app_dir / "brief.json"
        map_path = app_dir / "map.json"

    plan_path = app_dir / "plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    template = plan_module.load_template(args.workspace, cfg, args.kind)
    templates = cfg["templates"]
    sub = templates["resume"] if args.kind == "resume" else templates["cover_letter"]
    template_docx = Path(args.workspace) / templates["dir"] / sub / "template.docx"
    if not template_docx.is_file():
        raise CareerDocsError(f"template not found at {template_docx}", code="USAGE")

    # Baselines land directly under baselines/<positioning>/; role outputs under outputs/.
    out_dir = app_dir if args.baseline else app_dir / "outputs"
    org = ""
    if brief_path and brief_path.is_file():
        org = json.loads(brief_path.read_text(encoding="utf-8")).get("organization", "")
    stem = document_stem(plan, template, args.kind, cfg["outputs"]["file_name"],
                         slug=f"baseline-{positioning}" if args.baseline else slug, org=org)
    out_dir.mkdir(parents=True, exist_ok=True)
    rotate_previous(out_dir, stem)
    out_docx = out_dir / f"{stem}.docx"
    render_document(plan, template, template_docx, out_docx)
    linkify(out_docx, link_map(load_provider(args.workspace, cfg).read()))

    pdf_path = None
    pdf_available = shutil.which("soffice") is not None
    if args.pdf:
        pdf_path = convert_to_pdf(out_docx)
        pdf_available = pdf_path is not None

    record = build_record(
        plan, document=out_docx, kind=args.kind, positioning=plan["positioning"],
        plan_path=plan_path,
        brief_path=brief_path if (brief_path and brief_path.is_file()) else None,
        map_path=map_path if (map_path and map_path.is_file()) else None,
        pdf_path=pdf_path, pdf_available=pdf_available,
    )
    invalid = schema.validate_against("output-record", record)
    if invalid:
        raise CareerDocsError("output record is invalid: " + "; ".join(invalid))
    record_path = out_docx.with_suffix(".record.json")
    record_path.write_text(json.dumps(record, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    if args.json:
        print(json.dumps({"document": str(out_docx), "pdf": str(pdf_path) if pdf_path else None, "record": str(record_path)}))
    else:
        print(f"rendered {out_docx}")
        if args.pdf:
            print(f"pdf: {pdf_path}" if pdf_path else "pdf: skipped (soffice not available)")
    return 0
