#!/usr/bin/env python3
"""Build the binary example fixtures deterministically.

The fictional applicant "Jordan Rivera" is sanitized (``example.com`` email, ``555``
phone) and lives entirely under ``examples/`` and ``tests/fixtures/``. This script builds
the artifacts that cannot be hand-authored — a DOCX resume (python-docx) and a PDF resume
(reportlab) — with fixed content and fixed metadata so re-running produces the same
documents. The two resumes deliberately disagree on one end date (the Globex role), which
is the conflict the onboarding flow must surface. Text fixtures (config, voice, notes,
CSV, job description, extracted candidates) are committed directly, not generated here.

Run: ``uv run --extra dev python scripts/build_fixtures.py`` (``--templates-only`` rebuilds
just the two example templates).
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / "examples" / "applicant" / "sources"
TEMPLATES = ROOT / "examples" / "applicant" / "templates"
RENDERED = ROOT / "tests" / "fixtures" / "rendered"

FIXED_TIME = datetime(2024, 1, 1, 0, 0, 0)

CONTACT = {
    "name": "Jordan Rivera",
    "headline": "Engineering Leader",
    "location": "Metropolis, USA",
    "email": "jordan.rivera@example.com",
    "phone": "(555) 555-0142",
}

# The Globex end date is the deliberate conflict: the DOCX says 2021-06, the PDF 2021-08.
GLOBEX_END_DOCX = "June 2021"
GLOBEX_END_PDF = "August 2021"

ACHIEVEMENTS = [
    "Grew the platform engineering team from 4 to 15 engineers.",
    "Cut cloud infrastructure costs by 35% through workload consolidation.",
]
SKILLS = ["Python", "Go", "Kubernetes", "Team Leadership"]


def build_resume_a_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading(CONTACT["name"], level=0)
    document.add_paragraph(f"{CONTACT['headline']} — {CONTACT['location']}")
    document.add_paragraph(f"{CONTACT['email']} · {CONTACT['phone']}")

    document.add_heading("Experience", level=1)
    document.add_heading(f"Engineering Manager, Globex Corporation", level=2)
    document.add_paragraph(f"March 2018 – {GLOBEX_END_DOCX}")
    for achievement in ACHIEVEMENTS:
        document.add_paragraph(achievement, style="List Bullet")

    document.add_heading("Senior Software Engineer, Initech", level=2)
    document.add_paragraph("January 2015 – February 2018")
    document.add_paragraph("Built and operated the billing platform.", style="List Bullet")

    document.add_heading("Skills", level=1)
    document.add_paragraph(", ".join(SKILLS))

    document.add_heading("Education", level=1)
    document.add_paragraph("B.S. Computer Science, State University (2011 – 2015)")

    props = document.core_properties
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.author = CONTACT["name"]

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def build_resume_b_pdf(path: Path) -> None:
    from reportlab import rl_config

    rl_config.invariant = 1  # deterministic timestamps and ids
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    width, height = letter
    y = height - 72

    def line(text: str, size: int = 11, gap: int = 16) -> None:
        nonlocal y
        pdf.setFont("Helvetica", size)
        pdf.drawString(72, y, text)
        y -= gap

    line(CONTACT["name"], size=18, gap=24)
    line(f"{CONTACT['headline']} — {CONTACT['location']}")
    line(f"{CONTACT['email']} · {CONTACT['phone']}", gap=24)

    line("Experience", size=14, gap=20)
    line("Engineering Manager, Globex Corporation")
    line(f"March 2018 – {GLOBEX_END_PDF}", gap=20)
    line("Senior Software Engineer, Initech")
    line("January 2015 – February 2018", gap=24)

    line("Skills", size=14, gap=20)
    line(", ".join(SKILLS))

    pdf.setTitle("Jordan Rivera — Resume")
    pdf.setAuthor(CONTACT["name"])
    pdf.showPage()
    pdf.save()


RESUME_TEMPLATE_JSON = {
    "name": "example-resume",
    "kind": "resume",
    "version": "2",
    "page_budget": 2,
    "units_per_page": 22,
    "sections": [
        {"id": "header", "title": "", "placeholder": "contact", "entity_types": ["contact"], "max_items": 1, "required": True},
        {"id": "summary", "title": "Summary", "placeholder": "summary", "kind": "sentence", "entity_types": ["experience", "achievement", "project", "skill"], "max_items": 1, "required": False},
        {"id": "experience", "title": "Experience", "placeholder": "experience", "entity_types": ["experience", "achievement", "project"], "experience_kinds": ["employment"], "max_items": 30, "required": True, "role_summaries": True},
        {"id": "projects", "title": "Projects", "placeholder": "projects", "entity_types": ["project", "achievement"], "max_items": 8, "required": False},
        {"id": "education", "title": "Education", "placeholder": "education", "entity_types": ["education"], "max_items": 5, "required": False},
        {"id": "credentials", "title": "Certifications", "placeholder": "credentials", "entity_types": ["credential"], "max_items": 6, "required": False},
        {"id": "patents", "title": "Patents", "placeholder": "patents", "entity_types": ["patent"], "max_items": 6, "required": False},
        {"id": "publications", "title": "Publications", "placeholder": "publications", "entity_types": ["publication"], "max_items": 6, "required": False},
        {"id": "awards", "title": "Awards", "placeholder": "awards", "entity_types": ["award"], "max_items": 6, "required": False},
        {"id": "affiliations", "title": "Affiliations", "placeholder": "affiliations", "entity_types": ["affiliation", "experience", "achievement"], "experience_kinds": ["advising", "board"], "max_items": 8, "required": False},
        {"id": "volunteer", "title": "Volunteer Experience", "placeholder": "volunteer", "entity_types": ["experience", "achievement"], "experience_kinds": ["volunteer"], "max_items": 6, "required": False, "role_summaries": True},
        {"id": "skills", "title": "Technical Focus", "placeholder": "skills", "kind": "labeled", "group_by": "category", "entity_types": ["skill"], "max_items": 24, "required": False},
        {"id": "interests", "title": "Interests", "placeholder": "interests", "join": ", ", "entity_types": ["interest"], "max_items": 8, "required": False},
    ],
    "allowlist": ["Summary", "Experience", "Projects", "Education", "Certifications", "Patents", "Publications", "Awards", "Affiliations", "Volunteer Experience", "Technical Focus", "Interests"],
    "style_notes": "Two-page budget; verb-first bullets; no buzzwords. Single column, Calibri, 0.7-inch margins: centered name, ruled capitalized headings, bold role with the organization and location muted and the dates on a right tab, an italic descriptor under the role, projects as italic sub-heads under their role, bulleted achievements, bold-label skill lines, degrees and awards with the year on the right tab. Every section is optional except the header and Experience; the approach's `sections` picks the ones a role uses. Role descriptors go on the summary line, never in the role line.",
}

COVER_LETTER_TEMPLATE_JSON = {
    "name": "example-cover-letter",
    "kind": "cover_letter",
    "version": "2",
    "page_budget": 1,
    "units_per_page": 6,
    "sections": [
        {"id": "header", "title": "", "placeholder": "contact", "entity_types": ["contact"], "max_items": 1, "required": True},
        {"id": "body", "title": "", "placeholder": "body", "kind": "sentence", "entity_types": ["experience", "skill"], "max_items": 6, "required": True},
    ],
    "allowlist": ["Sincerely,"],
    "style_notes": "One page; complements the resume; draft prose in voice, never paste resume bullets. The greeting is the first body unit; the template places the role line and the date above it and carries the sign-off.",
}


_INK = "1F1F1F"
_MUTED = "555555"
_RULE = "999999"
_MARGIN_INCHES = 0.7
_TEXT_WIDTH_INCHES = 8.5 - 2 * _MARGIN_INCHES
_BODY_PT = 11
_SMALL_PT = 10.5
_LINE_SPACING = 1.05


def _base_document():
    """Letter page, 0.7-inch margins, Calibri 11 pt body in near-black."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Inches(_MARGIN_INCHES))
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(_BODY_PT)
    normal.font.color.rgb = RGBColor.from_string(_INK)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = _LINE_SPACING
    bullets = document.styles["List Bullet"]
    bullets.font.name, bullets.font.size = "Calibri", Pt(_BODY_PT)
    bullets.font.color.rgb = RGBColor.from_string(_INK)
    return document


def _run(paragraph, text, *, size=None, bold=None, italic=None, color=None, caps=False):
    from docx.shared import Pt, RGBColor

    run = paragraph.add_run(text)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if caps:
        run.font.all_caps = True  # displayed in capitals; the text itself stays as written
    return run


def _paragraph(document, text=None, *, size=None, bold=None, italic=None, color=None, center=False,
               before=None, after=None, style=None, caps=False, keep_next=False, indent=None,
               right_tab=False, line_spacing=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt

    paragraph = document.add_paragraph(style=style)
    fmt = paragraph.paragraph_format
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    if keep_next:
        fmt.keep_with_next = True
    if indent is not None:
        fmt.left_indent = Inches(indent)
    if line_spacing is not None:
        fmt.line_spacing = line_spacing
    if right_tab:
        fmt.tab_stops.add_tab_stop(Inches(_TEXT_WIDTH_INCHES), WD_TAB_ALIGNMENT.RIGHT)
    if text is not None:
        _run(paragraph, text, size=size, bold=bold, italic=italic, color=color, caps=caps)
    return paragraph


def _rule_below(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    border = OxmlElement("w:bottom")
    for attribute, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "2"), ("w:color", _RULE)):
        border.set(qn(attribute), value)
    borders = OxmlElement("w:pBdr")
    borders.append(border)
    paragraph._p.get_or_add_pPr().append(borders)


def _preserve_spaces(document) -> None:
    # Rendered values may begin or end with a space; keep every text run intact.
    from docx.oxml.ns import qn

    for text in document.element.body.iter(qn("w:t")):
        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _header(document, *, after: float) -> None:
    _paragraph(document, "{{ contact_name }}", size=20, bold=True, center=True, after=2)
    _paragraph(document, "{{ contact_details }}", size=_SMALL_PT, center=True, after=after)


def _tag(document, text: str) -> None:
    _paragraph(document, text)


def _finish(document, directory: Path, manifest: dict) -> None:
    _preserve_spaces(document)
    props = document.core_properties
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    directory.mkdir(parents=True, exist_ok=True)
    document.save(str(directory / "template.docx"))
    (directory / "template.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")


def _dated_line(document, *, before: float, after: float, lead_color=_INK, rest_color=_INK,
                tail_color=_INK, keep_next=False):
    """``lead`` bold, ``rest`` after it, ``tail`` on the right tab stop."""
    paragraph = _paragraph(document, before=before, after=after, keep_next=keep_next, right_tab=True)
    _run(paragraph, "{{ unit.lead }}", bold=True, color=lead_color)
    _run(paragraph, "{{ unit.rest }}", color=rest_color)
    paragraph.add_run().add_tab()
    _run(paragraph, "{{ unit.tail }}", color=tail_color)
    return paragraph


def build_resume_template(directory: Path) -> None:
    """Single column: centered name and contact line, ruled uppercase section headings, a
    bold role with the organization muted and the dates on a right tab stop, the role's
    summary beneath in italics, projects as italic sub-heads under their role, bulleted
    achievements, bold-label skill lines, and dated lines for degrees, awards, and the like."""
    document = _base_document()
    _header(document, after=3)
    _tag(document, "{%p for section in sections %}")
    _tag(document, "{%p if section.units %}")
    _rule_below(_paragraph(document, "{{ section.title }}", size=_BODY_PT, bold=True, before=10, after=3.5,
                           caps=True, keep_next=True))
    _tag(document, "{%p for unit in section.units %}")
    _tag(document, "{%p if unit.kind == 'bullet' %}")
    _paragraph(document, "{{ unit.text }}", style="List Bullet", after=2, line_spacing=1.08)
    _tag(document, "{%p elif unit.kind == 'subhead' %}")
    subhead = _paragraph(document, before=3.5, after=1, keep_next=True, indent=0.1)
    _run(subhead, "{{ unit.lead }}", bold=True, italic=True, color=_INK)
    _run(subhead, "{{ unit.rest }}", italic=True, color=_MUTED)
    _tag(document, "{%p elif unit.kind == 'labeled' and unit.tail %}")
    labeled = _paragraph(document, before=2)
    _run(labeled, "{{ unit.head }}", bold=True)
    _run(labeled, "  {{ unit.tail }}")
    _tag(document, "{%p elif unit.kind == 'field' and unit.tail and section.id == 'experience' %}")
    _dated_line(document, before=5.5, after=0, rest_color=_MUTED, tail_color=_MUTED, keep_next=True)
    _tag(document, "{%p if unit.note %}")
    _paragraph(document, "{{ unit.note }}", size=_SMALL_PT, italic=True, color=_MUTED, after=1)
    _tag(document, "{%p endif %}")
    _tag(document, "{%p elif unit.kind == 'field' and unit.tail %}")
    _dated_line(document, before=2, after=0, tail_color=_MUTED, keep_next=True)
    _tag(document, "{%p if unit.note %}")
    _paragraph(document, "{{ unit.note }}", size=_SMALL_PT, italic=True, color=_MUTED, after=1)
    _tag(document, "{%p endif %}")
    _tag(document, "{%p else %}")
    _paragraph(document, "{{ unit.text }}", after=3.5)
    _tag(document, "{%p endif %}")
    _tag(document, "{%p endfor %}")
    _tag(document, "{%p endif %}")
    _tag(document, "{%p endfor %}")
    _finish(document, directory, RESUME_TEMPLATE_JSON)


def build_cover_letter_template(directory: Path) -> None:
    """The same header, then the role line and the date, the body paragraphs, and a
    sign-off followed by the name."""
    document = _base_document()
    _header(document, after=5)
    _tag(document, "{%p if role_line %}")
    _paragraph(document, "{{ role_line }}", size=_BODY_PT, bold=True, center=True, after=22)
    _tag(document, "{%p endif %}")
    _paragraph(document, "{{ date_line }}", size=_BODY_PT, after=12)
    _tag(document, "{%p for section in sections %}")
    _tag(document, "{%p for unit in section.units %}")
    _paragraph(document, "{{ unit.text }}", size=_BODY_PT, after=9)
    _tag(document, "{%p endfor %}")
    _tag(document, "{%p endfor %}")
    _paragraph(document, "Sincerely,", size=_BODY_PT, before=7, after=16)
    _paragraph(document, "{{ contact_name }}", size=_BODY_PT)
    _finish(document, directory, COVER_LETTER_TEMPLATE_JSON)


_RESUME_LINES = [
    ("Jordan Rivera", 18, 26),
    ("Engineering Leader — Metropolis, USA", 11, 16),
    ("Metropolis, USA · (555) 555-0142 · jordan.rivera@example.com", 11, 24),
    ("Experience", 14, 20),
    ("Engineering Manager, Globex Corporation (2018-03–2021-06)", 11, 16),
    ("Grew the platform engineering team from 4 to 15 engineers.", 11, 16),
    ("Cut cloud infrastructure costs by 35% through workload consolidation.", 11, 24),
    ("Senior Software Engineer, Initech (2015-01–2018-02)", 11, 24),
    ("Skills", 14, 20),
    ("Python, Go, Kubernetes, Team Leadership", 11, 24),
    ("Education", 14, 20),
    ("B.S. Computer Science, State University", 11, 16),
]


def build_example_resume_pdf(path: Path) -> None:
    """A representative rendered resume PDF for the layout checks, with 1-inch margins."""
    from reportlab import rl_config

    rl_config.invariant = 1
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    _width, height = letter
    y = height - 72  # 1-inch top margin
    for text, size, gap in _RESUME_LINES:
        pdf.setFont("Helvetica-Bold" if size >= 14 else "Helvetica", size)
        pdf.drawString(72, y, text)  # 1-inch left margin
        y -= gap
    pdf.setTitle("Jordan Rivera — Resume")
    pdf.showPage()
    pdf.save()


def main(argv: list[str] | None = None) -> None:
    templates_only = "--templates-only" in (argv if argv is not None else sys.argv[1:])
    if not templates_only:
        build_resume_a_docx(SOURCES / "resume-a.docx")
        build_resume_b_pdf(SOURCES / "resume-b.pdf")
        build_example_resume_pdf(RENDERED / "example-resume.pdf")
        print(f"built {SOURCES / 'resume-a.docx'}")
        print(f"built {SOURCES / 'resume-b.pdf'}")
        print(f"built {RENDERED / 'example-resume.pdf'}")
    build_resume_template(TEMPLATES / "resume")
    build_cover_letter_template(TEMPLATES / "cover-letter")
    print(f"built {TEMPLATES / 'resume' / 'template.docx'}")
    print(f"built {TEMPLATES / 'cover-letter' / 'template.docx'}")


if __name__ == "__main__":
    main()
