"""Tests for docxtpl rendering and the output record."""

import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "skills" / "careerdocs" / "scripts"
sys.path.insert(0, str(SCRIPTS))

from careerdocs import brief, diff, mapping, merge, plan, render  # noqa: E402
from careerdocs.config import default_config  # noqa: E402
from careerdocs.providers import load_provider  # noqa: E402

CANDIDATES = json.loads((ROOT / "tests" / "fixtures" / "candidates.json").read_text())
JD = (ROOT / "examples" / "applicant" / "applications" / "example-role" / "job-description.md").read_text()
TEMPLATE_DIR = ROOT / "examples" / "applicant" / "templates" / "resume"
TEMPLATE_DOCX = TEMPLATE_DIR / "template.docx"
TEMPLATE_JSON = json.loads((TEMPLATE_DIR / "template.json").read_text())


def build_plan(tmp_path, positioning="builder"):
    provider = load_provider(tmp_path, default_config())
    ops = merge.build_operations(provider.read(), CANDIDATES)
    d = diff.make_diff(provider, ops)
    diff.approve(provider, d["diff_id"])
    diff.apply(provider, d["diff_id"], cfg=default_config(), workspace=tmp_path)
    b = brief.generate_brief(JD)
    m = mapping.generate_map(b, provider.read())
    return plan.generate_plan(m, provider.read(), TEMPLATE_JSON, positioning)


def read_docx_text(path):
    from docx import Document

    return "\n".join(p.text for p in Document(str(path)).paragraphs)


def test_render_produces_docx_with_plan_content(tmp_path):
    p = build_plan(tmp_path)
    out = tmp_path / "resume.docx"
    render.render_document(p, TEMPLATE_JSON, TEMPLATE_DOCX, out)
    text = read_docx_text(out)
    assert "Globex Corporation" in text
    assert "Cut deploy time" in text or "Grew the platform" in text
    assert "Python" in text


def test_render_omits_gap_requirement_content(tmp_path):
    p = build_plan(tmp_path)
    out = tmp_path / "resume.docx"
    render.render_document(p, TEMPLATE_JSON, TEMPLATE_DOCX, out)
    text = read_docx_text(out).lower()
    assert "fda" not in text
    assert "medical device" not in text


def test_document_stem_from_contact_and_kind(tmp_path):
    p = build_plan(tmp_path)
    assert render.slugify("Jordan Q. Rivera-Smith") == "Jordan-Q-Rivera-Smith"
    assert render.document_stem(p, TEMPLATE_JSON, "resume", "{name}-{kind}") == "Jordan-Rivera-Resume"
    assert render.document_stem(p, TEMPLATE_JSON, "cover_letter", "{name}-{kind}-{org}", org="Globex Corp") == "Jordan-Rivera-Cover-Letter-Globex-Corp"
    # An empty placeholder never leaves a dangling hyphen.
    assert render.document_stem(p, TEMPLATE_JSON, "resume", "{name}-{kind}-{org}") == "Jordan-Rivera-Resume"


def test_unit_context_splits_role_dates_and_note():
    unit = {"text": "Eng, Acme\tMar 2018 – Jun 2021\nA fintech startup.", "kind": "field"}
    context = render._unit_context(unit)
    assert (context["head"], context["tail"], context["note"]) == ("Eng, Acme", "Mar 2018 – Jun 2021", "A fintech startup.")


def test_build_record_shape(tmp_path):
    p = build_plan(tmp_path)
    record = render.build_record(
        p, document=Path("out.docx"), kind="resume", positioning="builder",
        plan_path=Path("plan.json"), brief_path=Path("brief.json"), map_path=Path("map.json"),
        pdf_path=None, pdf_available=False,
    )
    expected_sources = {sid for u in p["units"] for sid in u["source_ids"]}
    assert set(record["source_ids"]) == expected_sources
    assert set(record["checks"]) == {"factual", "links_dates", "extraction", "pagination", "layout"}
    # Without a PDF, pagination and layout are skipped.
    assert record["checks"]["pagination"]["status"] == "skipped"
    assert record["checks"]["factual"]["status"] == "pending"
    assert record["stale"] is False


def test_cli_render_writes_docx_and_record(tmp_path, capsys):
    from careerdocs import cli

    ws = str(tmp_path)
    cli.main(["config", "init", "--workspace", ws])
    # Provision the template in the workspace.
    dest = tmp_path / "templates" / "resume"
    dest.mkdir(parents=True)
    shutil.copy(TEMPLATE_DOCX, dest / "template.docx")
    shutil.copy(TEMPLATE_DIR / "template.json", dest / "template.json")
    # Build the pipeline up to the plan.
    p = build_plan(tmp_path)
    app = tmp_path / "applications" / "example-role"
    app.mkdir(parents=True)
    (app / "plan.json").write_text(json.dumps(p), encoding="utf-8")

    capsys.readouterr()
    rc = cli.main(["render", "--kind", "resume", "--role-slug", "example-role", "--workspace", ws, "--json"])
    assert rc == 0
    out = json.loads(capsys.readouterr().out)
    doc = Path(out["document"])
    assert doc.exists() and doc.name == "Jordan-Rivera-Resume.docx"
    record = json.loads(Path(out["record"]).read_text())
    assert record["kind"] == "resume" and record["source_ids"]


def test_render_rotates_the_previous_render(tmp_path, capsys):
    from careerdocs import cli

    ws = str(tmp_path)
    cli.main(["config", "init", "--workspace", ws])
    dest = tmp_path / "templates" / "resume"
    dest.mkdir(parents=True)
    shutil.copy(TEMPLATE_DOCX, dest / "template.docx")
    shutil.copy(TEMPLATE_DIR / "template.json", dest / "template.json")
    p = build_plan(tmp_path)
    app = tmp_path / "applications" / "example-role"
    app.mkdir(parents=True)
    (app / "plan.json").write_text(json.dumps(p), encoding="utf-8")
    outputs = app / "outputs"
    argv = ["render", "--kind", "resume", "--role-slug", "example-role", "--workspace", ws, "--json"]

    assert cli.main(argv) == 0
    (outputs / "layout").mkdir()
    (outputs / "layout" / "Jordan-Rivera-Resume-p1.png").write_bytes(b"png")
    first = (outputs / "Jordan-Rivera-Resume.docx").read_bytes()
    capsys.readouterr()

    assert cli.main(argv) == 0
    assert cli.main(argv) == 0
    names = sorted(f.name for f in outputs.iterdir() if f.is_file())
    assert names == sorted([
        "Jordan-Rivera-Resume.docx", "Jordan-Rivera-Resume.record.json",
        "Jordan-Rivera-Resume_bak1.docx", "Jordan-Rivera-Resume_bak1.record.json",
        "Jordan-Rivera-Resume_bak2.docx", "Jordan-Rivera-Resume_bak2.record.json",
    ])
    # The oldest render is now _bak2, its layout render moved with it, and its record
    # points at its new path; the newest carries the plain name.
    assert (outputs / "Jordan-Rivera-Resume_bak2.docx").read_bytes() == first
    assert (outputs / "layout" / "Jordan-Rivera-Resume_bak2-p1.png").exists()
    record = json.loads((outputs / "Jordan-Rivera-Resume_bak2.record.json").read_text())
    assert record["document"] == str(outputs / "Jordan-Rivera-Resume_bak2.docx")


COVER_DIR = ROOT / "examples" / "applicant" / "templates" / "cover-letter"
COVER_JSON = json.loads((COVER_DIR / "template.json").read_text())


def test_render_cover_letter(tmp_path):
    provider = load_provider(tmp_path, default_config())
    ops = merge.build_operations(provider.read(), CANDIDATES)
    d = diff.make_diff(provider, ops)
    diff.approve(provider, d["diff_id"])
    diff.apply(provider, d["diff_id"], cfg=default_config(), workspace=tmp_path)
    b = brief.generate_brief(JD)
    m = mapping.generate_map(b, provider.read())
    p = plan.generate_plan(m, provider.read(), COVER_JSON, "builder", brief=b)
    out = tmp_path / "letter.docx"
    render.render_document(p, COVER_JSON, COVER_DIR / "template.docx", out)
    text = read_docx_text(out)
    assert "Jordan Rivera" in text  # contact header
    # The body carries evidence content.
    assert "Globex Corporation" in text or "Kubernetes" in text or "platform" in text


def test_cli_baseline_plan_render_check(tmp_path, capsys):
    from careerdocs import cli

    ws = str(tmp_path)
    cli.main(["config", "init", "--workspace", ws])
    dest = tmp_path / "templates" / "resume"
    dest.mkdir(parents=True)
    shutil.copy(TEMPLATE_DOCX, dest / "template.docx")
    shutil.copy(TEMPLATE_DIR / "template.json", dest / "template.json")
    # Build a profile in the workspace.
    provider = load_provider(tmp_path, default_config())
    ops = merge.build_operations(provider.read(), CANDIDATES)
    d = diff.make_diff(provider, ops)
    diff.approve(provider, d["diff_id"])
    diff.apply(provider, d["diff_id"], cfg=default_config(), workspace=tmp_path)

    assert cli.main(["plan", "--baseline", "--positioning", "builder", "--workspace", ws]) == 0
    assert (tmp_path / "baselines" / "builder" / "plan.json").exists()

    capsys.readouterr()
    assert cli.main(["render", "--baseline", "--positioning", "builder", "--workspace", ws, "--json"]) == 0
    out = json.loads(capsys.readouterr().out)
    doc = Path(out["document"])
    assert doc.parent == tmp_path / "baselines" / "builder"
    record = json.loads(Path(out["record"]).read_text())
    assert record["role_brief"] is None  # baseline has no role

    capsys.readouterr()
    assert cli.main(["check", str(doc), "--workspace", ws, "--json"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["checks"]["factual"]["status"] == "pass"


def test_cli_render_pdf_branch(tmp_path):
    from careerdocs import cli

    ws = str(tmp_path)
    cli.main(["config", "init", "--workspace", ws])
    dest = tmp_path / "templates" / "resume"
    dest.mkdir(parents=True)
    shutil.copy(TEMPLATE_DOCX, dest / "template.docx")
    shutil.copy(TEMPLATE_DIR / "template.json", dest / "template.json")
    p = build_plan(tmp_path)
    app = tmp_path / "applications" / "example-role"
    app.mkdir(parents=True)
    (app / "plan.json").write_text(json.dumps(p), encoding="utf-8")

    rc = cli.main(["render", "--kind", "resume", "--role-slug", "example-role", "--pdf", "--workspace", ws])
    assert rc == 0
    assert (app / "outputs" / "Jordan-Rivera-Resume.docx").exists()
    if shutil.which("soffice"):
        assert (app / "outputs" / "Jordan-Rivera-Resume.pdf").exists()


def test_build_context_exposes_contact_lines_and_unit_halves(tmp_path):
    p = build_plan(tmp_path)
    context = render.build_context(p, TEMPLATE_JSON)
    assert context["contact_name"] == "Jordan Rivera"
    assert context["contact"] == f"Jordan Rivera\n{context['contact_details']}"
    assert "jordan.rivera@example.com" in context["contact_details"]
    experience = next(s for s in context["sections"] if s["id"] == "experience")
    role = next(u for u in experience["units"] if u["kind"] == "field")
    assert "\t" not in role["head"] and role["tail"]  # role, then dates


def test_render_puts_the_name_on_its_own_line(tmp_path):
    p = build_plan(tmp_path)
    out = tmp_path / "resume.docx"
    render.render_document(p, TEMPLATE_JSON, TEMPLATE_DOCX, out)
    lines = read_docx_text(out).splitlines()
    assert lines[0] == "Jordan Rivera"
    assert "jordan.rivera@example.com" in lines[1]


def test_link_map_uses_bare_display_forms():
    profile = {"entities": [
        {"type": "contact", "email": "a@example.com", "visibility": "public", "verification": "applicant_verified",
         "links": [{"label": "LinkedIn", "url": "https://www.linkedin.com/in/x/"}]},
        {"type": "patent", "url": "https://patents.example.com/p1", "visibility": "public", "verification": "applicant_verified"},
        {"type": "project", "links": [{"label": "Repo", "url": "https://github.com/x/y"}], "visibility": "private", "verification": "applicant_verified"},
    ]}
    assert render.link_map(profile) == {
        "linkedin.com/in/x": "https://www.linkedin.com/in/x/",
        "patents.example.com/p1": "https://patents.example.com/p1",
        "a@example.com": "mailto:a@example.com",
    }


def test_linkify_makes_profile_links_clickable_without_changing_text(tmp_path):
    p = build_plan(tmp_path)
    profile = load_provider(tmp_path, default_config()).read()
    out = tmp_path / "resume.docx"
    render.render_document(p, TEMPLATE_JSON, TEMPLATE_DOCX, out)
    before = read_docx_text(out)
    assert render.linkify(out, render.link_map(profile)) >= 1  # at least the email
    assert read_docx_text(out) == before  # the visible text is unchanged, so the checks still hold
    from docx import Document

    document = Document(str(out))
    targets = {rel.target_ref for rel in document.part.rels.values() if rel.reltype.endswith("/hyperlink")}
    assert "mailto:jordan.rivera@example.com" in targets
    assert document.element.body.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink")
